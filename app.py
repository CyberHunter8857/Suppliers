from flask import Flask, render_template, request, send_file
from docx import Document
from pymongo import MongoClient
import io
import datetime

app = Flask(__name__)

# MongoDB Atlas connection setup
client = MongoClient("mongodb+srv://admin:<db_password>@cyberhunter.h1bcf.mongodb.net/?retryWrites=true&w=majority&appName=CyberHunter")
db = client["assignmentDB"]
collection = db["submissions"]

# Function to replace placeholders with student details in the docx file
def replace_placeholders_in_docx(name, class_name, roll_number):
    doc = Document("assignment.docx")
    
    for paragraph in doc.paragraphs:
        if 'xyz' in paragraph.text:
            paragraph.text = paragraph.text.replace('xyz', name)
        if 'classssss' in paragraph.text:
            paragraph.text = paragraph.text.replace('classssss', class_name)
        if '11111111' in paragraph.text:
            paragraph.text = paragraph.text.replace('11111111', roll_number)

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    
    return doc_io

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate', methods=['POST'])
def generate():
    student_name = request.form['student_name']
    class_name = request.form['class_name']
    roll_number = request.form['roll_number']

    # Save to MongoDB
    submission = {
        "name": student_name,
        "class": class_name,
        "roll_number": roll_number,
        "timestamp": datetime.datetime.now()
    }
    collection.insert_one(submission)
    
    # Generate modified document
    updated_docx = replace_placeholders_in_docx(student_name, class_name, roll_number)
    
    return send_file(
        updated_docx,
        as_attachment=True,
        download_name=f"assignment_{student_name}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

if __name__ == '__main__':
    app.run(debug=True)
